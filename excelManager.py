import openpyxl, os
from openpyxl import load_workbook
from openpyxl.styles import Alignment, PatternFill, Border, Side, Font


def insertInfoRunAndLegende(nome_file, dict_info_Excel, dict_periferiche) :

    lista_ip_periferiche_Excel = []
    for key in dict_periferiche :    
        if key == "barriera_ottica": nome_completo = "Barriera Ottica" 
        if key == "antenna": nome_completo = "Antenna RFID"
        if key == "anpr_zona_1": nome_completo = "Telecamera ANPR Zona 1"
        if key == "anpr_zona_2": nome_completo = "Telecamera ANPR Zona 2"
        if key == "anpr_zona_3": nome_completo = "Telecamera ANPR Zona 3"
        if key == "bio": nome_completo = "BIO Board e DAC Driver"
        if key == "stampante": nome_completo = "Stampante"
        if key == "fare_display": nome_completo = "Fare Display"
        if key == "WIM" : nome_completo = "WIM"          
        tmp = f'{nome_completo}: {dict_periferiche[key]["ip"]}' 
        lista_ip_periferiche_Excel.append(tmp)
    
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Parametri Tool"
    bordi = Border(left=Side(style='thin'),
                     right=Side(style='thin'),
                     top=Side(style='thin'),
                     bottom=Side(style='thin'))
    bold_font = Font(bold=True)
    centering = Alignment(horizontal="center", vertical="center")

    ws.merge_cells('A1:B1')
    header_cell = ws.cell(row=1, column=1, value="Parametri Tool")
    header_cell.font = Font(bold=True)
    header_cell.alignment = centering
   
    ws.insert_rows(2)

    ws["A2"] = "Parametro"
    ws["B2"] = "Valore"
    ws["A2"].font = bold_font
    ws["B2"].font = bold_font
    ws["A2"].border = bordi
    ws["B2"].border = bordi
    ws["A2"].alignment = centering
    ws["B2"].alignment = centering

    row_i = 3
    col_widths = [len("Parametro"), len("Valore")]

    dict_info_Excel["lista_ip_periferiche"] = lista_ip_periferiche_Excel

    for k, val in dict_info_Excel.items():
        if isinstance(val, list):  
            for elem in val:
                k_lista, val_lista = elem.split(": ")
                ws.append([k_lista, val_lista])
                ws[f"A{row_i}"].border = bordi
                ws[f"B{row_i}"].border = bordi
                ws[f"A{row_i}"].alignment = centering
                ws[f"B{row_i}"].alignment = centering
                col_widths[0] = max(col_widths[0], len(k_lista))
                col_widths[1] = max(col_widths[1], len(val_lista))
                ws.row_dimensions[row_i].height = 25
                row_i += 1
        else:
            ws.append([k, val])
            ws[f"A{row_i}"].border = bordi
            ws[f"B{row_i}"].border = bordi
            ws[f"A{row_i}"].alignment = centering
            ws[f"B{row_i}"].alignment = centering
            col_widths[0] = max(col_widths[0], len(k))
            col_widths[1] = max(col_widths[1], len(str(val)))
            ws.row_dimensions[row_i].height = 25 
            row_i += 1

    ws.column_dimensions["A"].width = 30
    ws.column_dimensions["B"].width = 20
    ws.row_dimensions[2].height = 25

    
    
    # Creo la Prima legenda...
    colore_rosso = PatternFill(start_color="FF6666", end_color="FF6666", fill_type="solid")
    colore_verde = PatternFill(start_color="66CC66", end_color="66CC66", fill_type="solid")
    colore_giallo = PatternFill(start_color="FFEE99", end_color="FFEE99", fill_type="solid")
    colore_arancione = PatternFill(start_color="e87c1b", end_color="e87c1b", fill_type="solid")

    data = [
        ("Esito", "Descrizione output ottenuto"),
        ("Yes", "La macchina o periferica è stata raggiunta"),
        ("Not Expected", "La macchina o periferica non necessita di essere raggiunta"),
        ("No", "La macchina o periferica non è stata raggiunta") ]

    ws.row_dimensions[1].height = 25
    ws.merge_cells('D1:E1')
    header_cell = ws.cell(row=1, column=4, value="Legenda 'Raggiungibilità Piste Egitto'")
    header_cell.font = Font(bold=True)
    header_cell.alignment = centering


    for row_idx, row_data in enumerate(data, start=2):
        for col_idx, value in enumerate(row_data, start=4): 
            row = ws.cell(row=row_idx, column=col_idx, value=value)
            row.alignment = Alignment(horizontal="left", vertical="center")
            row.border = bordi

    
    for row_idx in range(3, 6):  
        color_cell = ws.cell(row=row_idx, column=4)
        if color_cell.value == "Yes":
            color_cell.fill = colore_verde
        elif color_cell.value == "No":
            color_cell.fill = colore_rosso
        elif color_cell.value == "Not Expected":
            color_cell.fill = colore_giallo
        row = ws.cell(row=row_idx, column=4)
        row.alignment = centering

    # Creo la Seconda legenda...
    data = [
        ("Esito", "Descrizione output ottenuto"),
        ("<version>", "La versione Software è stata recuperata con successo"),
        ("Not Expected", "Il recupero della versione Software non è necessario"),
        ("<WARNING>", "Sono state recuperate due versioni Software differenti"),
        ("Unreachable", "La macchina è irraggiungibile, non è stato possibile recuperare la versione Software"),
        ("Error", "È stato riscontrato un errore durante il recupero della versione Software") ]

    ws.merge_cells('D7:E7')
    header_cell = ws.cell(row=7, column=4, value="Legenda 'Verifica Versioni Software'")
    header_cell.font = Font(bold=True)
    header_cell.alignment = centering

    for row_idx, row_data in enumerate(data, start=8):
        for col_idx, value in enumerate(row_data, start=4): 
            row = ws.cell(row=row_idx, column=col_idx, value=value)
            row.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
            row.border = bordi

    
    for row_idx in range(9,14):  
        color_cell = ws.cell(row=row_idx, column=4)
        if color_cell.value == "<version>":
            color_cell.fill = colore_verde
        elif color_cell.value == "<WARNING>":
            color_cell.fill = colore_arancione
            row = ws.cell(row=row_idx, column=4, value="<version>")
        elif color_cell.value == "Unreachable":
            color_cell.fill = colore_rosso
        elif color_cell.value == "Not Expected":
            color_cell.fill = colore_giallo
        elif color_cell.value == "Error":
            color_cell.fill = colore_rosso
        row = ws.cell(row=row_idx, column=4)
        row.alignment = centering

    ws["D2"].alignment = centering
    ws["E2"].alignment = centering

    ws["D8"].alignment = centering
    ws["E8"].alignment = centering
    
    # Creo la Terza legenda...
    colore_rosso = PatternFill(start_color="FF6666", end_color="FF6666", fill_type="solid")
    colore_verde = PatternFill(start_color="66CC66", end_color="66CC66", fill_type="solid")
    colore_giallo = PatternFill(start_color="FFEE99", end_color="FFEE99", fill_type="solid")

    data = [
        ("Esito", "Descrizione output ottenuto (tolleranza = -10 minuti rispetto sysdate)"),
        ("OK", "Tutti i file in '/acme/spool/transits/' e '../photos/' rispettano la tolleranza"),
        ("Transits KO", "Uno o più file presenti in '/acme/spool/transits/' non rispettano la tolleranza"),
        ("Photos KO", "Uno o più file presenti in '/acme/spool/photos/' non rispettano la tolleranza"),
        ("KO", "Alcuni file presenti sia in '/acme/spool/transits/' che in '../photos/' non rispettano la tolleranza"),
        ("Unreachable", "Non è stato possibile verificare la vecchiaia dei file causa irraggiungibilità del Controller") ]

    
    ws.merge_cells('D15:E15')
    header_cell = ws.cell(row=15, column=4, value="Legenda 'Acquisizione file acme spool'")
    header_cell.font = Font(bold=True)
    header_cell.alignment = centering


    for row_idx, row_data in enumerate(data, start=16):
        for col_idx, value in enumerate(row_data, start=4): 
            row = ws.cell(row=row_idx, column=col_idx, value=value)
            row.alignment = Alignment(horizontal="left", vertical="center")
            row.border = bordi

    for row_idx in range(15,22):
        ws.row_dimensions[row_idx].height = 25  
        color_cell = ws.cell(row=row_idx, column=4)
        if color_cell.value == "OK":
            color_cell.fill = colore_verde
        elif color_cell.value == "KO":
            color_cell.fill = colore_rosso
        elif color_cell.value == "Transits KO":
            color_cell.fill = colore_giallo
        elif color_cell.value == "Photos KO":
            color_cell.fill = colore_giallo
        elif color_cell.value == "Unreachable":
            color_cell.fill = colore_rosso
        row = ws.cell(row=row_idx, column=4)
        row.alignment = centering

    
    ws.column_dimensions["D"].width = 15
    ws.column_dimensions["E"].width = 85

    ws["D16"].alignment = centering
    ws["E16"].alignment = centering

    wb.save(nome_file)

def insertPingResultsIntoExcel(nome_file, dict_all_piste, colonne_Excel_ping) :
    
    wb = load_workbook(nome_file)
    ws = wb.create_sheet("Raggiungibilità Piste Egitto")
    bold_colonne = Font(bold=True)

    bordi = Border(left=Side(border_style="thin", color="000000"), 
                   right=Side(border_style="thin", color="000000"), 
                   top=Side(border_style="thin", color="000000"), 
                   bottom=Side(border_style="thin", color="000000"))

    colore_rosso = PatternFill(start_color="FF6666", end_color="FF6666", fill_type="solid")
    colore_verde = PatternFill(start_color="66CC66", end_color="66CC66", fill_type="solid")
    colore_giallo = PatternFill(start_color="FFEE99", end_color="FFEE99", fill_type="solid")
    colore_neutro = PatternFill(start_color="FFFFFF", end_color="FFFFFF", fill_type="solid")

    for i_col, nome_colonna in enumerate(colonne_Excel_ping, start=1) :
        colonna = ws.cell(row=1, column=i_col, value=nome_colonna)
        col_letter = ws.cell(row=1, column=i_col).column_letter  
        if nome_colonna == "WIM":
            nome_fittizio = "NotExpected"
            ws.column_dimensions[col_letter].width = max(len(nome_fittizio) + 2, 10)
        else:
            ws.column_dimensions[col_letter].width = max(len(nome_colonna) + 2, 10)

        colonna.alignment = Alignment(horizontal="center", vertical="center")
        colonna.border = bordi
        colonna.font = bold_colonne
    
    for i, key in enumerate(dict_all_piste) :
        exc_codice_pista = str(dict_all_piste[key]["codice_pista"])
        exc_tipo_pista = str(dict_all_piste[key]["tipo_pista"])
        exc_controller_pista = str(dict_all_piste[key]["controller_pista"][0]["stato"])
        exc_barriera_ottica = str(dict_all_piste[key]["controller_pista"][1]["barriera_ottica"]["stato"])
        exc_antenna = str(dict_all_piste[key]["controller_pista"][1]["antenna"]["stato"])
        exc_anpr_zona_1 = str(dict_all_piste[key]["controller_pista"][1]["anpr_zona_1"]["stato"])
        exc_anpr_zona_2 = str(dict_all_piste[key]["controller_pista"][1]["anpr_zona_2"]["stato"])
        exc_anpr_zona_3 = str(dict_all_piste[key]["controller_pista"][1]["anpr_zona_3"]["stato"])
        exc_bio = str(dict_all_piste[key]["controller_pista"][1]["bio"]["stato"])
        exc_stampante = str(dict_all_piste[key]["controller_pista"][1]["stampante"]["stato"])
        exc_fare_display = str(dict_all_piste[key]["controller_pista"][1]["fare_display"]["stato"])
        exc_WIM = str(dict_all_piste[key]["controller_pista"][1]["WIM"]["stato"])
        exc_external = str(dict_all_piste[key]["axis_external"]['stato'])
        exc_internal = str(dict_all_piste[key]["axis_internal"]['stato'])

        dati_pista = (exc_codice_pista, exc_tipo_pista, exc_controller_pista, 
                        exc_barriera_ottica, exc_antenna, exc_anpr_zona_1, exc_anpr_zona_2, exc_anpr_zona_3,exc_bio, exc_stampante, exc_fare_display, exc_WIM, 
                        exc_external, exc_internal)

        for col_i, element in enumerate(dati_pista, start=1):  
            colore = colore_neutro
            match element :
                case "T":
                    element = "Yes"
                    colore = colore_verde
                case "F":
                    element = "No"
                    colore = colore_rosso
                case "N":
                    element = "Not Expected"
                    colore = colore_giallo

            row = ws.cell(row=i+2, column=col_i, value=element)
            row.alignment = Alignment(horizontal="center", vertical="center")
            row.fill = colore
            row.border = bordi
    
    for i in range(1, ws.max_row + 1):
        ws.row_dimensions[i].height = 19

    wb.save(nome_file)
    
def insertVersionResultsIntoExcel(nome_file_Excel, dict_all_piste, colonne_Excel_versioni) :

    wb = load_workbook(nome_file_Excel)
    ws2 = wb.create_sheet("Verifica Versioni Software")


    bold_colonne = Font(bold=True)
    bordi = Border(left=Side(border_style="thin", color="000000"), 
                   right=Side(border_style="thin", color="000000"), 
                   top=Side(border_style="thin", color="000000"), 
                   bottom=Side(border_style="thin", color="000000"))

    colore_rosso = PatternFill(start_color="FF6666", end_color="FF6666", fill_type="solid")
    colore_giallo = PatternFill(start_color="FFEE99", end_color="FFEE99", fill_type="solid")
    colore_verde = PatternFill(start_color="66CC66", end_color="90EE90", fill_type="solid")
    colore_arancione = PatternFill(start_color="e87c1b", end_color="e87c1b", fill_type="solid")

    for i_col, nome_colonna in enumerate(colonne_Excel_versioni, start=1) :
        colonna = ws2.cell(row=1, column=i_col, value=nome_colonna)
        col_letter = ws2.cell(row=1, column=i_col).column_letter  
        colonna.alignment = Alignment(horizontal="center", vertical="center")
        colonna.border = bordi
        colonna.font = bold_colonne
    
    for i, key in enumerate(dict_all_piste) :
        
        exc_codice_pista = str(dict_all_piste[key]["codice_pista"])
        vrs_controller_pista = str(dict_all_piste[key]["controller_pista"][0]["versione"])
        
        vrs_pacc_movyonEgypt = str(dict_all_piste[key]["controller_pista"][0]["versione_pacchetti"]["movyonEgypt-"])
        vrs_pacc_movyonEgyptConf = str(dict_all_piste[key]["controller_pista"][0]["versione_pacchetti"]["movyonEgyptConf"])
        vrs_pacc_arcd_application_base = str(dict_all_piste[key]["controller_pista"][0]["versione_pacchetti"]["arcd_application_base"])
        vrs_pacc_arcdLibs = str(dict_all_piste[key]["controller_pista"][0]["versione_pacchetti"]["arcdLibs"])
        vrs_pacc_arcd = str(dict_all_piste[key]["controller_pista"][0]["versione_pacchetti"]["arcd-"])
        vrs_pacc_arcd_application_core = str(dict_all_piste[key]["controller_pista"][0]["versione_pacchetti"]["arcd_application_core-"])
        
        vrs_barriera_ottica = str(dict_all_piste[key]["controller_pista"][1]["barriera_ottica"]["versione"])
        vrs_antenna = str(dict_all_piste[key]["controller_pista"][1]["antenna"]["versione"])
        vrs_anpr_zona_1 = str(dict_all_piste[key]["controller_pista"][1]["anpr_zona_1"]["versione"])
        vrs_anpr_zona_2 = str(dict_all_piste[key]["controller_pista"][1]["anpr_zona_2"]["versione"])
        vrs_anpr_zona_3 = str(dict_all_piste[key]["controller_pista"][1]["anpr_zona_3"]["versione"])
        vrs_bio = str(dict_all_piste[key]["controller_pista"][1]["bio"]["versioneBIO"])
        vrs_dac = str(dict_all_piste[key]["controller_pista"][1]["bio"]["versioneDAC"])
          
        vrs_external = str(dict_all_piste[key]["axis_external"]['versione'])
        vrs_internal = str(dict_all_piste[key]["axis_internal"]['versione'])
    
        versioni_software_pista = (exc_codice_pista, vrs_controller_pista, 
                                vrs_pacc_movyonEgypt, vrs_pacc_movyonEgyptConf, vrs_pacc_arcd_application_base, vrs_pacc_arcdLibs, vrs_pacc_arcd, vrs_pacc_arcd_application_core,
                                vrs_barriera_ottica, vrs_antenna, vrs_anpr_zona_1, vrs_anpr_zona_2, vrs_anpr_zona_3, vrs_bio, vrs_dac, vrs_external, vrs_internal)
    
        for col_i, element in enumerate(versioni_software_pista, start=1):  
                colore = colore_verde
                match element :
                    case "Not Expected":
                        colore = colore_giallo
                    case "Unreachable":
                        colore = colore_rosso
                    case "Errore":
                        element = "Error"
                        colore = colore_rosso
                
                if " -> (WARNING)" in element :
                    colore = colore_arancione
                    element = element.replace(" -> (WARNING)", "")
                row = ws2.cell(row=i+2, column=col_i, value=element)
                row.alignment = Alignment(horizontal="center", vertical="center")
            
                if col_i != 1:
                    row.fill = colore
                row.border = bordi
    
        for col in ws2.columns:
            max_length = 0
            col_letter = col[0].column_letter 
            for cell in col:
                if cell.value:
                    max_length = max(max_length, len(str(cell.value)))  
            ws2.column_dimensions[col_letter].width = max_length + 2 
  
        for i in range(1, ws2.max_row + 1):
            ws2.row_dimensions[i].height = 19

    wb.save(nome_file_Excel)

def insertStatoTransitiFotoIntoExcel(nome_file_Excel, dict_all_piste) :
    wb = load_workbook(nome_file_Excel)
    ws = wb.create_sheet("Acquisizione file acme spool")
    bordi = Border(left=Side(style='thin'),
                     right=Side(style='thin'),
                     top=Side(style='thin'),
                     bottom=Side(style='thin'))
    bold_font = Font(bold=True)
    centering = Alignment(horizontal="center", vertical="center")

    colore_rosso = PatternFill(start_color="FF6666", end_color="FF6666", fill_type="solid")
    colore_verde = PatternFill(start_color="66CC66", end_color="66CC66", fill_type="solid")
    colore_giallo = PatternFill(start_color="FFEE99", end_color="FFEE99", fill_type="solid")

    ws["A1"] = "IP Controller"
    ws["B1"] = "Transits&Photos"
    ws["A1"].font = bold_font
    ws["B1"].font = bold_font
    ws["A1"].border = bordi
    ws["B1"].border = bordi
    ws["A1"].alignment = centering
    ws["B1"].alignment = centering

    for i,key in enumerate(dict_all_piste) :
        ip_controller = dict_all_piste[key]["controller_pista"][0]["ip"]
        stato = dict_all_piste[key]["controller_pista"][0]["stato_transiti_foto"]
        
        row = ws.cell(row=i+2, column=1, value=ip_controller)
        row.alignment = Alignment(horizontal="center", vertical="center")
        row.border = bordi
        row = ws.cell(row=i+2, column=2, value=stato)
        
        match stato :
                case "OK":
                    colore = colore_verde
                case "KO":
                    colore = colore_rosso
                case "Transits KO" | "Photos KO":
                    colore = colore_giallo
                case "Unreachable":
                    colore = colore_rosso
        row.fill = colore
        row.border = bordi
        row.alignment = Alignment(horizontal="center", vertical="center")
        
    ws.column_dimensions["A"].width = 18
    ws.column_dimensions["B"].width = 18
    
    for row in range(1, ws.max_row + 1):  
        ws.row_dimensions[row].height = 25


    wb.save(nome_file_Excel)

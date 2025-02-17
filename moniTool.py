import dbManager
import sys, os, time, math, re, time
from paramiko import AutoAddPolicy, ProxyCommand, SSHClient, Channel
from datetime import datetime
import openpyxl
from openpyxl import load_workbook
from openpyxl.styles import Alignment, PatternFill, Border, Side, Font


'''------------------------------------- PROVE -------------------------------------'''

    





#sys.exit(1)


def checkPingMachine(ip) -> str :  
    res = os.popen(f"ping {ip}").read()
    if ("Trasmessi = 4") in res and ("Ricevuti = 0") not in res :
        if ("TTL scaduto durante il passaggio.") in res:
            print(f"Macchina (ip: {ip}) non raggiungibile!\n")
            return "F"
        print(f"Macchina (ip: {ip}) raggiunta!\n")
        return "T"
    else :
        print(f"Macchina (ip: {ip}) non raggiungibile!\n")
        return "F"

def checkPingPeriferiche(dict_periferiche, ssh_client, tipo_pista) -> str :
    
    lista_stato_periferiche='' 
    
    for key in dict_periferiche :
        
        if key == "anpr_zona_2" and tipo_pista != 'I' :
            print(f"---> Periferica anpr_zona_2 non presa in considerazione per pista di tipo {tipo_pista}")
            lista_stato_periferiche+="N"
        elif key == "anpr_zona_3" and tipo_pista in ('H','I') :
             print(f"---> Periferica anpr_zona_3 non presa in considerazione per pista di tipo {tipo_pista}")
             lista_stato_periferiche+="N"
        elif key == "WIM" and tipo_pista != 'K' :
             print(f"---> Periferica WIM non presa in considerazione per pista di tipo {tipo_pista}")
             lista_stato_periferiche+="N"
        else :          
            ip_periferica = dict_periferiche[key]["ip"]
            stdin, stdout, stderr = ssh_client.exec_command(f'ping -c 4 {ip_periferica}')
            res = stdout.read().decode()

            if ("4 packets transmitted,") in res and ("0 received,") not in res :    
                if("Time to live exceeded") in res :
                    print(f"---> Periferica {key} (ip: {ip_periferica}) non raggiungibile causa TTL scaduta!")
                    lista_stato_periferiche+="F"
                else :
                    print(f"---> Periferica {key} (ip: {ip_periferica}) raggiunta!")
                    lista_stato_periferiche+="T"    
            else :
                print(f"---> Periferica {key} (ip: {ip_periferica}) non raggiungibile")
                lista_stato_periferiche+="F"
        
    return lista_stato_periferiche

    print("CIAO")

def getInfoBySSHSession(controller_pista, dict_pacchetti_controller, dict_periferiche, user, pwd, tipo_pista) : 
    ssh_client = SSHClient()
    ssh_client.set_missing_host_key_policy(AutoAddPolicy())
    
    lista_stato_periferiche=''
    
    try:       
        ssh_client.connect(hostname=controller_pista, username=user, password=pwd, allow_agent=True, look_for_keys=True)
        print(f"---> Avvio connessione in ssh ({controller_pista})\n")
        
        print("Recupero versioni controller di Pista e pacchetti relativi...")
        vrs_ctr, lista_vrs_pacchetti = getVersioneControllerAndPacchetti(controller_pista, dict_pacchetti_controller, user, pwd, ssh_client)
        print("Versioni controller e pacchetti recuperate con successo!\n")
        
        print("Verifico raggiungibilità delle periferiche di Pista...")
        lista_stato_periferiche = checkPingPeriferiche(dict_periferiche, ssh_client, tipo_pista)
        print("Verifica raggiungibilità periferiche terminata!\n")

        print("Recupero versioni SW delle periferiche di Pista...")
        dict_versioni_periferiche = getVersioniPeriferiche(ssh_client, dict_periferiche, lista_stato_periferiche)
        print("Versioni SW recuperate con successo!\n")

    finally : 
        
        print(f"---> Chiudo connessione in ssh ({controller_pista})\n")
        
        ssh_client.close()
        #sys.exit(1) # TO DO TOGLIERE
        return vrs_ctr, lista_vrs_pacchetti, lista_stato_periferiche, dict_versioni_periferiche

def getVersioneControllerAndPacchetti (controller_pista, dict_pacchetti_controller, user, pwd, ssh_client: SSHClient) -> str :
    
    versioni_pacchetti=[]
    stdin, stdout, stderr = ssh_client.exec_command('cat /etc/ansible/facts.d/customfacts.fact')
    vrs_ctr = stdout.read().decode().strip()
    vrs_ctr = re.search(r"=\s*(.*)", vrs_ctr)
    vrs_ctr = vrs_ctr.group(1)

    for key in dict_pacchetti_controller :
        stdin, stdout, stderr = ssh_client.exec_command(f'rpm -qa | grep {key}')
        vrs_pac = stdout.read().decode().strip()
        vrs_pac = vrs_pac.replace("\n", ", ")
        versioni_pacchetti.append(vrs_pac)

        print(f"---> Versione pacchetto {key} recuperata: {vrs_pac}") 

    return vrs_ctr, versioni_pacchetti

def getVersioniPeriferiche(ssh_client_pista: SSHClient, dict_periferiche, lista_stato_periferiche) -> dict : 
    
    dict_versioni_periferiche = {"v_barriera_ottica" : '', 
                                "v_antenna" : '', 
                                "v_anpr_zona_1" : '', 
                                "v_anpr_zona_2" : '', 
                                "v_anpr_zona_3" : '', 
                                "v_bio": ['', '']}
    
    unreachable_vrs = "Unreachable"

    for i, key in enumerate(dict_periferiche) :
        stato_periferica = lista_stato_periferiche[i]
        ip_periferica = dict_periferiche[key]["ip"]

        match key:
            case "barriera_ottica" :
                
                if stato_periferica == "F" :
                    print(f"---Versione Barriera Ottica : {unreachable_vrs}")
                    dict_versioni_periferiche["v_barriera_ottica"] = unreachable_vrs
                else :
                    user_barriera = "root"
                    pwd_barriera = "controlcard"
                    stringa_versione = "Mif version is: "
                    barriera_cmd = "/var/mif/mif -v"
                                        
                    cmd = f"sshpass -p {pwd_barriera} ssh {user_barriera}@{ip_periferica} '{barriera_cmd}' "
                    stdin, stdout, stderr = ssh_client_pista.exec_command(cmd)
                    file_content = stdout.read().decode()
                    versione_barriera = file_content.split(f"{stringa_versione}")[1].split("\n")[0]
                
                    print(f"---> Versione Barriera Ottica recuperata: {versione_barriera}")
                    dict_versioni_periferiche["v_barriera_ottica"] = versione_barriera
                
            case "antenna" :
                
                if stato_periferica == "F" :
                    print(f"---> Versione Antenna : {unreachable_vrs}")
                    dict_versioni_periferiche["v_antenna"] = unreachable_vrs
                else :
                    user_antenna = "root"
                    pwd_antenna = "UHF-RFID-Dev"
                    stringa_versione = "KAT ACQ"     
                    antenna_cmd = "grep 'KAT ACQ' /flash/log/KatAcqFFM.log"
                    
                    cmd = f"sshpass -p {pwd_antenna} ssh {user_antenna}@{ip_periferica} '{antenna_cmd}' "
                    stdin, stdout, stderr = ssh_client_pista.exec_command(cmd)
                    file_content = stdout.read().decode()
                    lines = file_content.split("\n")

                    last_line = None
                    for line in lines:
                        if stringa_versione in line:
                            last_line = line 
                    m = re.search(rf"--\s*({stringa_versione} \d+\.\d+)\s*--", last_line)
                    versione_antenna = m.group(1)
                    
                    print(f"---> Versione Antenna recuperata: {versione_antenna}") 
                    dict_versioni_periferiche["v_antenna"] = versione_antenna

            case "anpr_zona_1" if stato_periferica == "T":
                
                user_anpr_zona_1 = ""
                pwd_anpr_zona_1 = ""
                versione_anp_zona_1 = ""
 
                if stato_periferica == "F" :
                    print(f"---> Versione anpr_zona_1 : {unreachable_vrs}")
                    dict_versioni_periferiche["v_anpr_zona_1"] = unreachable_vrs
                else :
                    print("---> Versione anpr_zona_1: momentaneamente non disponibile")
                    dict_versioni_periferiche["v_anpr_zona_1"] = "Non disponibile"
            
            case "anpr_zona_2" :
                
                user_anpr_zona_2 = ""
                pwd_anpr_zona_2 = ""
                versione_anp_zona_2 = ""
 
                if stato_periferica == "F" :
                    print(f"---> Versione anpr_zona_2 : {unreachable_vrs}")
                    dict_versioni_periferiche["v_anpr_zona_2"] = unreachable_vrs
                else :
                    print("---> Versione anpr_zona_2: momentaneamente non disponibile")
                    dict_versioni_periferiche["v_anpr_zona_2"] = "Non disponibile"
            
            case "anpr_zona_3" :
                
                user_anpr_zona_3 = ""
                pwd_anpr_zona_3 = ""
                versione_anp_zona_3 = ""
 
                if stato_periferica == "F" :
                    print(f"---> Versione anpr_zona_3 : {unreachable_vrs}")
                    dict_versioni_periferiche["v_anpr_zona_3"] = unreachable_vrs
                else :
                    print("---> Versione anpr_zona_3: momentaneamente non disponibile")
                    dict_versioni_periferiche["v_anpr_zona_3"] = "Non disponibile"
            
            case "bio" :                
                
                if stato_periferica == "F" :
                    print(f"---> Versione DAC : {unreachable_vrs}") 
                    print(f"---> Versione BIO : {unreachable_vrs}")
                    dict_versioni_periferiche["v_bio"][0] = unreachable_vrs
                    dict_versioni_periferiche["v_bio"][1] = unreachable_vrs
                else :
                    user_bio = "root"
                    pwd_bio = "itsegypt"  
                    dac_cmd = "/FN0/S30DacDRIVER/dac -v"
                    bio_cmd = "cat /FN0/version.txt"
                    versione_BIO = ''
                    stringhe_versioni_bio = ["BIO_EGYPT v", "BIO v", "SAM_ATECH v"]
                                        
                    cmd = f"sshpass -p {pwd_bio} ssh {user_bio}@{ip_periferica} '{dac_cmd}' "
                    stdin, stdout, stderr = ssh_client_pista.exec_command(cmd)
                    file_content = stdout.read().decode()
                    versione_DAC = file_content.strip()
                    
                    cmd = f"sshpass -p {pwd_bio} ssh {user_bio}@{ip_periferica} '{bio_cmd}' "
                    stdin, stdout, stderr = ssh_client_pista.exec_command(cmd)
                    file_content = stdout.read().decode()
                    
                    for k in stringhe_versioni_bio:
                        match = re.search(rf"{k}\d+\.\d+", file_content)
                        versione_BIO += match.group() + ", "              
                    versione_BIO = versione_BIO.rstrip(', ')

                    print(f"---> Versione DAC recuperata: {versione_DAC}") 
                    print(f"---> Versione BIO recuperata: {versione_BIO}")
                    dict_versioni_periferiche["v_bio"][0] = versione_DAC
                    dict_versioni_periferiche["v_bio"][1] = versione_BIO
    
    return dict_versioni_periferiche

def insertPingResultsIntoExcel(nome_file, dict_all_piste, colonne_Excel_ping) :
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Raggiungibilità_piste_Egitto"    
    
    bold_colonne = Font(bold=True)

    bordi = Border(left=Side(border_style="thin", color="000000"), right=Side(border_style="thin", color="000000"), top=Side(border_style="thin", color="000000"), bottom=Side(border_style="thin", color="000000"))

    colore_rosso = PatternFill(start_color="FF6666", end_color="FF6666", fill_type="solid")
    colore_verde = PatternFill(start_color="66CC66", end_color="66CC66", fill_type="solid")
    colore_giallo = PatternFill(start_color="FFEE99", end_color="FFEE99", fill_type="solid")
    colore_neutro = PatternFill(start_color="FFFFFF", end_color="FFFFFF", fill_type="solid")

    for i_col, nome_colonna in enumerate(colonne_Excel_ping, start=1) :
        colonna = ws.cell(row=1, column=i_col, value=nome_colonna)
        col_letter = ws.cell(row=1, column=i_col).column_letter  
        if nome_colonna == "WIM":
            nome_fittizio = "NotExpected"
            ws.column_dimensions[col_letter].width = max(len(nome_fittizio) + 2, 10)
        else:
            ws.column_dimensions[col_letter].width = max(len(nome_colonna) + 2, 10)

        colonna.alignment = Alignment(horizontal="center", vertical="center")
        colonna.border = bordi
        colonna.font = bold_colonne
    
    for i, key in enumerate(dict_all_piste) :
        exc_codice_pista = str(dict_all_piste[key]["codice_pista"])
        exc_tipo_pista = str(dict_all_piste[key]["tipo_pista"])
        exc_controller_pista = str(dict_all_piste[key]["controller_pista"][0]["stato"])
        exc_barriera_ottica = str(dict_all_piste[key]["controller_pista"][1]["barriera_ottica"]["stato"])
        exc_antenna = str(dict_all_piste[key]["controller_pista"][1]["antenna"]["stato"])
        exc_anpr_zona_1 = str(dict_all_piste[key]["controller_pista"][1]["anpr_zona_1"]["stato"])
        exc_anpr_zona_2 = str(dict_all_piste[key]["controller_pista"][1]["anpr_zona_2"]["stato"])
        exc_anpr_zona_3 = str(dict_all_piste[key]["controller_pista"][1]["anpr_zona_3"]["stato"])
        exc_bio = str(dict_all_piste[key]["controller_pista"][1]["bio"]["stato"])
        exc_stampante = str(dict_all_piste[key]["controller_pista"][1]["stampante"]["stato"])
        exc_fare_display = str(dict_all_piste[key]["controller_pista"][1]["fare_display"]["stato"])
        exc_WIM = str(dict_all_piste[key]["controller_pista"][1]["WIM"]["stato"])
        exc_external = str(dict_all_piste[key]["axis_external"]['stato'])
        exc_internal = str(dict_all_piste[key]["axis_internal"]['stato'])

        dati_pista = (exc_codice_pista, exc_tipo_pista, exc_controller_pista, 
                        exc_barriera_ottica, exc_antenna, exc_anpr_zona_1, exc_anpr_zona_2, exc_anpr_zona_3,exc_bio, exc_stampante, exc_fare_display, exc_WIM, 
                        exc_external, exc_internal)

        for col_i, element in enumerate(dati_pista, start=1):  
            colore = colore_neutro
            match element :
                case "T":
                    element = "Yes"
                    colore = colore_verde
                case "F":
                    element = "No"
                    colore = colore_rosso
                case "N":
                    element = "Not Expected"
                    colore = colore_giallo

            row = ws.cell(row=i+2, column=col_i, value=element)
            row.alignment = Alignment(horizontal="center", vertical="center")
            row.fill = colore
            row.border = bordi
    
    for i in range(1, ws.max_row + 1):
        ws.row_dimensions[i].height = 19

    wb.save(nome_file)
    
def insertVersionResultsIntoExcel(nome_file_Excel, dict_all_piste, colonne_Excel_versioni) :
    wb = load_workbook(nome_file_Excel)
    ws2 = wb.create_sheet("Versioni_Software")


    bold_colonne = Font(bold=True)
    bordi = Border(left=Side(border_style="thin", color="000000"), right=Side(border_style="thin", color="000000"), top=Side(border_style="thin", color="000000"), bottom=Side(border_style="thin", color="000000"))

    colore_rosso = PatternFill(start_color="FF6666", end_color="FF6666", fill_type="solid")
    colore_giallo = PatternFill(start_color="FFEE99", end_color="FFEE99", fill_type="solid")
    colore_verde = PatternFill(start_color="66CC66", end_color="90EE90", fill_type="solid")

    for i_col, nome_colonna in enumerate(colonne_Excel_versioni, start=1) :
        colonna = ws2.cell(row=1, column=i_col, value=nome_colonna)
        col_letter = ws2.cell(row=1, column=i_col).column_letter  
        colonna.alignment = Alignment(horizontal="center", vertical="center")
        colonna.border = bordi
        colonna.font = bold_colonne
    
    for i, key in enumerate(dict_all_piste) :
        
        exc_codice_pista = str(dict_all_piste[key]["codice_pista"])
        vrs_controller_pista = str(dict_all_piste[key]["controller_pista"][0]["versione"])
        
        vrs_pacc_movyonEgypt = str(dict_all_piste[key]["controller_pista"][0]["versione_pacchetti"]["movyonEgypt-"])
        vrs_pacc_movyonEgyptConf = str(dict_all_piste[key]["controller_pista"][0]["versione_pacchetti"]["movyonEgyptConf"])
        vrs_pacc_arcd_application_base = str(dict_all_piste[key]["controller_pista"][0]["versione_pacchetti"]["arcd_application_base"])
        vrs_pacc_arcdLibs = str(dict_all_piste[key]["controller_pista"][0]["versione_pacchetti"]["arcdLibs"])
        vrs_pacc_arcd = str(dict_all_piste[key]["controller_pista"][0]["versione_pacchetti"]["arcd-"])
        vrs_pacc_arcd_application_core = str(dict_all_piste[key]["controller_pista"][0]["versione_pacchetti"]["arcd_application_core-"])
        
        vrs_barriera_ottica = str(dict_all_piste[key]["controller_pista"][1]["barriera_ottica"]["versione"])
        vrs_antenna = str(dict_all_piste[key]["controller_pista"][1]["antenna"]["versione"])
        vrs_anpr_zona_1 = str(dict_all_piste[key]["controller_pista"][1]["anpr_zona_1"]["versione"])
        vrs_anpr_zona_2 = str(dict_all_piste[key]["controller_pista"][1]["anpr_zona_2"]["versione"])
        vrs_anpr_zona_3 = str(dict_all_piste[key]["controller_pista"][1]["anpr_zona_3"]["versione"])
        vrs_bio = str(dict_all_piste[key]["controller_pista"][1]["bio"]["versioneBIO"])
        vrs_dac = str(dict_all_piste[key]["controller_pista"][1]["bio"]["versioneDAC"])
          
        vrs_external = str(dict_all_piste[key]["axis_external"]['versione'])
        vrs_internal = str(dict_all_piste[key]["axis_internal"]['versione'])
    
        versioni_software_pista = (exc_codice_pista, vrs_controller_pista, 
                                vrs_pacc_movyonEgypt, vrs_pacc_movyonEgyptConf, vrs_pacc_arcd_application_base, vrs_pacc_arcdLibs, vrs_pacc_arcd, vrs_pacc_arcd_application_core,
                                vrs_barriera_ottica, vrs_antenna, vrs_anpr_zona_1, vrs_anpr_zona_2, vrs_anpr_zona_3, vrs_bio, vrs_dac, vrs_external, vrs_internal)
    
        for col_i, element in enumerate(versioni_software_pista, start=1):  
                colore = colore_verde
                match element :
                    case "Non disponibile":
                        colore = colore_giallo
                    case "Unreachable":
                        colore = colore_rosso

                row = ws2.cell(row=i+2, column=col_i, value=element)
                row.alignment = Alignment(horizontal="center", vertical="center")
            
                if col_i != 1:
                    row.fill = colore
                row.border = bordi
    
        for col in ws2.columns:
            max_length = 0
            col_letter = col[0].column_letter 
            for cell in col:
                if cell.value:
                    max_length = max(max_length, len(str(cell.value)))  
            ws2.column_dimensions[col_letter].width = max_length + 2 
  
        for i in range(1, ws2.max_row + 1):
            ws2.row_dimensions[i].height = 19

    wb.save(nome_file_Excel)


user="root"
pwd="Movyon.2022!"
controller_pista = '10.1.65.4'
inizio_esecuzione = datetime.now().replace(microsecond=0)

piste = dbManager.select_data_from_tjmd06_lne(dbManager.connessione_db_gsegy06())
#piste = (('1', 'K', '10.1.65.4', '10.1.123.132', '10.1.64.164'), ('2', 'I ', '10.1.65.8', '10.1.64.133', '10.1.123.165'))


dict_barriera_ottica = {"ip" : '172.16.0.10', "stato" : '', "versione" : ''}
dict_antenna = {"ip" : '172.16.0.30', "stato" : '', "versione" : ''}
dict_anpr_zona_1 = {"ip" : '172.16.0.21', "stato" : '', "versione" : ''}
dict_anpr_zona_2 = {"ip" : '172.16.0.22', "stato" : '', "versione" : ''}
dict_anpr_zona_3 = {"ip" : '172.16.0.22', "stato" : '', "versione" : ''}
dict_bio = {"ip" : '172.16.0.2', "stato" : '', "versioneBIO" : '', "versioneDAC" : ''}
dict_stampante = {"ip" : '172.16.0.40', "stato" : ''}
dict_fare_display = {"ip" : '172.16.0.60', "stato" : ''}
dict_WIM = {"ip" : '172.16.0.70', "stato" : ''}

dict_periferiche = {"barriera_ottica" : dict_barriera_ottica,
                   "antenna" : dict_antenna,
                   "anpr_zona_1" : dict_anpr_zona_1,
                    "anpr_zona_2" : dict_anpr_zona_2,
                    "anpr_zona_3" : dict_anpr_zona_3,
                    "bio" : dict_bio,
                    "stampante" : dict_stampante,
                    "fare_display" : dict_fare_display,
                    "WIM" : dict_WIM}

dict_pacchetti_controller = {"movyonEgypt-" :'',
                             "movyonEgyptConf" :'',
                             "arcd_application_base" :'',
                             "arcdLibs" :'',
                             "arcd-" :'',
                             "arcd_application_core-" :''}

dict_controller = {"ip" : '', "stato" : '', "versione" :'', "versione_pacchetti" :dict_pacchetti_controller}
dict_axis_external = {"ip" : '', "stato" : '', "versione" : ''}
dict_axis_internal = {"ip" : '', "stato" : '', "versione" : ''}


dict_pista = {"codice_pista" : '',
              "tipo_pista" : '',
              "controller_pista" : [dict_controller,  dict_periferiche], 
              "axis_external" : dict_axis_external,
              "axis_internal" : dict_axis_internal}

dict_all_piste = {}

print(f"\nInizio esecuzione: {inizio_esecuzione}")

for j in range(0, len(piste)) : 

    dict_pacchetti_controller = {"movyonEgypt-" :'', "movyonEgyptConf" :'', "arcd_application_base" :'', "arcdLibs" :'', "arcd-" :'', "arcd_application_core-" :''}
    dict_barriera_ottica = {"ip" : '172.16.0.10', "stato" : '', "versione" : ''}
    dict_antenna = {"ip" : '172.16.0.30', "stato" : '', "versione" : ''}
    dict_anpr_zona_1 = {"ip" : '172.16.0.21', "stato" : '', "versione" : ''}
    dict_anpr_zona_2 = {"ip" : '172.16.0.22', "stato" : '', "versione" : ''}
    dict_anpr_zona_3 = {"ip" : '172.16.0.22', "stato" : '', "versione" : ''}
    dict_bio = {"ip" : '172.16.0.2', "stato" : '', "versioneBIO" : '', "versioneDAC" : ''}
    dict_stampante = {"ip" : '172.16.0.40', "stato" : ''}
    dict_fare_display = {"ip" : '172.16.0.60', "stato" : ''}
    dict_WIM = {"ip" : '172.16.0.70', "stato" : ''}
    dict_periferiche = {"barriera_ottica" : dict_barriera_ottica, "antenna" : dict_antenna, "anpr_zona_1" : dict_anpr_zona_1, "anpr_zona_2" : dict_anpr_zona_2, "anpr_zona_3" : dict_anpr_zona_3, "bio" : dict_bio, "stampante" : dict_stampante, "fare_display" : dict_fare_display, "WIM" : dict_WIM}
    dict_controller = {"ip" : '', "stato" : '', "versione" : '', "versione_pacchetti" :dict_pacchetti_controller}
    dict_axis_external = {"ip" : '', "stato" : '', "versione" : ''}
    dict_axis_internal = {"ip" : '', "stato" : '', "versione" : ''}
    dict_pista = {"codice_pista" : '', "tipo_pista" : '', "controller_pista" : [dict_controller,  dict_periferiche], "axis_external" : dict_axis_external, "axis_internal" : dict_axis_internal}

    codice_pista = int(piste[j][0])
    tipo_pista = piste[j][1]
    print(f"\n------------------------ Inizio l'analisi della Pista_{codice_pista} ({j+1} di {len(piste)})------------------------\n")

    for i, key in enumerate(dict_pista) :

        info_pista = piste[j][i] 
        if key in ("axis_external", "axis_internal") :

            dict_pista[key]["ip"] = info_pista
            dict_pista[key]["stato"] = checkPingMachine(info_pista)
            dict_pista[key]["versione"] = "Non disponibile" # per il momento...

        elif key == "controller_pista" :
            
            controller_pista = info_pista
            dict_pista[key][0]["ip"] = info_pista
            dict_pista[key][0]["stato"] = checkPingMachine(info_pista)
            
            if dict_pista[key][0]["stato"] == "F" :
                print("Controller di Pista non raggiungibile, considero non raggiungibili le periferiche relative")
                lista_stato_periferiche = 'FFFFFFFFF'
                # TO DO chiamare la funzione per settare ad 'Unreachable' tutte le periferiche
                                
            else :
                versione_controller, lista_versione_pacchetti, lista_stato_periferiche, dict_versione_periferiche = getInfoBySSHSession(controller_pista, 
                                                                                                        dict_pacchetti_controller, 
                                                                                                        dict_periferiche, 
                                                                                                        user, pwd, tipo_pista)
                dict_pista[key][0]["versione"] = versione_controller.strip()

                for i, key_pacchetto in enumerate(dict_pacchetti_controller) :
                    dict_pista[key][0]["versione_pacchetti"][key_pacchetto] = lista_versione_pacchetti[i].strip()

                for x, key in enumerate(dict_periferiche) :
                    ip_periferica = dict_periferiche[key]["ip"]
                    dict_periferiche[key]["stato"] = lista_stato_periferiche[x]

                for key in dict_periferiche :
                    if f"v_{key}" in dict_versione_periferiche :
                        if key == "bio" :
                            dict_pista["controller_pista"][1][key]["versioneDAC"] = dict_versione_periferiche["v_bio"][0]
                            dict_pista["controller_pista"][1][key]["versioneBIO"] = dict_versione_periferiche["v_bio"][1]
                        else :     
                            dict_pista["controller_pista"][1][key]["versione"] = dict_versione_periferiche[f"v_{key}"]
        
        elif key == "codice_pista" :
            dict_pista[key] = int(info_pista)

        else :  
            dict_pista[key] = info_pista
    
    dict_all_piste["Pista_"+ str(codice_pista)] = dict_pista.copy()



#print(dict_all_piste)

nome_file_Excel="report_piste_Egitto.xlsx"

colonne_Excel_ping = ('Codice Pista', 'Tipo Pista', 'Controller', 'Barriera Ottica', 'Antenna RFID', 
                 'Telecamera ANPR Zona 1', 'Telecamera ANPR Zona 2', 'Telecamera ANPR Zona 3',
                  'BIO Board', 'Stampante', 'Fare Display', 'WIM', 'Telecamere CCTV Esterna ', 'Telecamere CCTV Interna ')

colonne_Excel_versioni = ('Codice Pista', 'Controller', 
                          'movyonEgypt', 'movyonEgyptConf', 'arcd_application_base', 'arcdLibs', 'arcd', 'arcd_application_core',
                          'Barriera Ottica', 'Antenna RFID', 
                          'Telecamera ANPR Zona 1', 'Telecamera ANPR Zona 2', 'Telecamera ANPR Zona 3',
                          'Bio Board', 'DAC Driver', 'Telecamera CCTV esterna', 'Telecamera CCTV interna')


insertPingResultsIntoExcel(nome_file_Excel, dict_all_piste, colonne_Excel_ping)
insertVersionResultsIntoExcel(nome_file_Excel, dict_all_piste, colonne_Excel_versioni)

print(f"\nFile '{nome_file_Excel}' generato con successo!")
print("File salvato in:", os.path.abspath(nome_file_Excel))


fine_esecuzione = datetime.now().replace(microsecond=0)
print(f"Fine esecuzione: {fine_esecuzione}")
delta = fine_esecuzione - inizio_esecuzione
delta_min = delta.total_seconds()/60
delta_min = math.trunc(delta_min * 100)/100

print(f"Tempo di esecuzione totale: {delta_min} minuti")


